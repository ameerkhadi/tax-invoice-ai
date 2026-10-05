"""Load invoices -> validate -> SQL model -> Excel report."""
from __future__ import annotations

import csv
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import DATA_DIR, OUT_DIR
from .validate import validate

SCHEMA = """
DROP TABLE IF EXISTS invoice; DROP TABLE IF EXISTS issue;
CREATE TABLE invoice (invoice_id TEXT PRIMARY KEY, supplier_id TEXT, supplier_name TEXT, supplier_org_no TEXT,
  invoice_date TEXT, month TEXT, category TEXT, vat_class TEXT, net_amount REAL, vat_amount REAL, total_amount REAL);
CREATE TABLE issue (invoice_id TEXT, rule TEXT, severity TEXT, detail TEXT);
CREATE VIEW IF NOT EXISTS vat_by_month AS
  SELECT month, vat_class, ROUND(SUM(net_amount),2) net, ROUND(SUM(vat_amount),2) vat, COUNT(*) n
  FROM invoice GROUP BY month, vat_class;
"""


def load_csv(path: str | None = None) -> list[dict]:
    with open(path or os.path.join(DATA_DIR, "invoices.csv"), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def build_db(rows: list[dict], db_path: str) -> list[tuple]:
    issues = validate(rows)
    con = sqlite3.connect(db_path)
    con.executescript(SCHEMA)
    con.executemany("INSERT OR IGNORE INTO invoice VALUES (?,?,?,?,?,?,?,?,?,?,?)", [
        (r["invoice_id"], r["supplier_id"], r["supplier_name"], r["supplier_org_no"], r["invoice_date"],
         r["invoice_date"][:7], r["category"], r["vat_class"], float(r["net_amount"]), float(r["vat_amount"]),
         float(r["total_amount"])) for r in rows])
    con.executemany("INSERT INTO issue VALUES (?,?,?,?)", issues)
    con.commit()
    con.close()
    return issues


def excel_report(db_path: str, xlsx_path: str) -> str:
    con = sqlite3.connect(db_path)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0B3C6E")

    def sheet(ws, title, q):
        ws.title = title
        cur = con.execute(q)
        cols = [d[0] for d in cur.description]
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for row in cur:
            ws.append(list(row))
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(col) + 6))
        ws.freeze_panes = "A2"

    sheet(wb.active, "Summary", """
        SELECT 'Invoices' metric, COUNT(*) value FROM invoice UNION ALL
        SELECT 'Net amount (SEK)', ROUND(SUM(net_amount),2) FROM invoice UNION ALL
        SELECT 'VAT amount (SEK)', ROUND(SUM(vat_amount),2) FROM invoice UNION ALL
        SELECT 'Invoices with issues', COUNT(DISTINCT invoice_id) FROM issue UNION ALL
        SELECT 'High-severity issues', COUNT(*) FROM issue WHERE severity='high'""")
    sheet(wb.create_sheet(), "Issues", "SELECT i.*, v.supplier_name, v.net_amount, v.vat_amount FROM issue i "
                                       "LEFT JOIN invoice v USING(invoice_id) ORDER BY severity, rule")
    sheet(wb.create_sheet(), "VAT by month", "SELECT * FROM vat_by_month ORDER BY month, vat_class")
    sheet(wb.create_sheet(), "Invoices", "SELECT * FROM invoice ORDER BY invoice_date")
    con.close()
    wb.save(xlsx_path)
    return xlsx_path


def run(db_path: str = os.path.join(OUT_DIR, "tax.db")) -> dict:
    os.makedirs(OUT_DIR, exist_ok=True)
    rows = load_csv()
    issues = build_db(rows, db_path)
    excel_report(db_path, os.path.join(OUT_DIR, "vat_review.xlsx"))
    from collections import Counter
    return {"invoices": len(rows), "issues": len(issues), "by_rule": dict(Counter(i[1] for i in issues))}


if __name__ == "__main__":
    print(run())
